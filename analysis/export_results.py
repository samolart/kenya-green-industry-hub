"""Export every number the deck needs from the recalculated Excel model to outputs/results.json.

The workbook is the single source of truth: the deck never contains a number that is not read from here.
Usage: python analysis/export_results.py model/kenya_green_hub_model.xlsx outputs/results.json
"""
import json
import sys
from openpyxl import load_workbook

src = sys.argv[1] if len(sys.argv) > 1 else "model/kenya_green_hub_model.xlsx"
dst = sys.argv[2] if len(sys.argv) > 2 else "outputs/results.json"
wb = load_workbook(src, data_only=True)
YEARS = list(range(2026, 2033))


def find_rows(ws, col=1, start=1):
    return {ws.cell(r, col).value: r for r in range(start, ws.max_row + 1) if ws.cell(r, col).value}


def series(ws, row, first_col=3, n=len(YEARS)):
    return [ws.cell(row, first_col + i).value for i in range(n)]


out = {"years": YEARS}

# ---- Grid ----
ws = wb["Grid"]
out["grid"] = {}
cur = None
for r in range(4, ws.max_row + 1):
    lab = ws.cell(r, 1).value
    if lab and str(lab).endswith("supply case"):
        cur = lab.split()[0]
        out["grid"][cur] = {}
        continue
    if cur is None or not lab:
        continue
    key = {
        "Peak demand": "peak",
        "Baseline firm available capacity (Jun 2026, incl. imports)": "baseline",
        "Firm evening-capable capacity": "firm",
        "Reserve margin before hub load": "rm",
        "Spare MW (firm capacity minus peak)": "spare",
        "New firm MW needed each year just to hold the margin": "need",
    }.get(lab.strip())
    if key is None and lab.strip().startswith("+"):
        key = "add_" + lab.strip()[2:].strip()
    if key:
        out["grid"][cur][key] = series(ws, r)

# ---- Hub ----
ws = wb["Hub"]
out["hub"] = {}
cur = None
for r in range(4, ws.max_row + 1):
    lab = ws.cell(r, 1).value
    if lab in ("Status quo", "Targeted hub", "Ambitious hub"):
        cur = lab
        out["hub"][cur] = {}
        continue
    if lab and lab.startswith("2030 prize"):
        cur = "_prize"
        out["hub"][cur] = {}
        continue
    if cur and cur != "_prize" and lab:
        key = {
            "Kenya data-centre IT capacity": "it",
            "IT capacity above Status quo": "incr",
            "Hub load on the grid at the evening peak": "load",
            "Infrastructure-fee revenue pool (Base fee)": "rev",
            "Cumulative build investment since 2026 (Base cost)": "capex",
        }.get(lab)
        if key:
            out["hub"][cur][key] = series(ws, r)
    elif cur == "_prize" and lab:
        out["hub"]["_prize"][lab] = {"Status quo": ws.cell(r, 3).value, "Targeted hub": ws.cell(r, 4).value, "Ambitious hub": ws.cell(r, 5).value}

# ---- Matrix ----
ws = wb["Matrix"]
out["matrix"] = {"rm_by_year": {}, "rm2030": {}, "gapmin": {}, "gaptgt": {}, "mwh": {}, "connect_min": {}, "connect_tgt": {}}
section = None
for r in range(5, ws.max_row + 1):
    a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
    if a in ("Stress", "Base", "Favourable") and b in ("Status quo", "Targeted hub", "Ambitious hub") and ws.cell(r, 3).value is not None and section is None:
        out["matrix"]["rm_by_year"][f"{a}|{b}"] = series(ws, r)
        continue
    if a and str(a).startswith("Reserve margin in 2030"):
        section = "rm2030"; continue
    if a and str(a).startswith("Firm MW needed in 2030 to reach the minimum"):
        section = "gapmin"; continue
    if a and str(a).startswith("Firm MW needed in 2030 to reach the target"):
        section = "gaptgt"; continue
    if a and str(a).startswith("Storage (MWh"):
        section = "mwh"; continue
    if a and str(a).startswith("Connectable new data-centre IT capacity (MW) while"):
        section = "connect_min"; continue
    if a and str(a).startswith("Connectable new data-centre IT capacity (MW) at the target"):
        section = "connect_tgt"; continue
    if section in ("rm2030", "gapmin", "gaptgt", "mwh") and b in ("Stress", "Base", "Favourable"):
        out["matrix"][section][b] = {"Status quo": ws.cell(r, 3).value, "Targeted hub": ws.cell(r, 4).value, "Ambitious hub": ws.cell(r, 5).value}
    if section in ("connect_min", "connect_tgt") and b in ("Stress", "Base", "Favourable"):
        out["matrix"][section][b] = series(ws, r)

# ---- UnitEconomics ----
ws = wb["UnitEconomics"]
u = {"tco": {}}
for r in range(5, 14):
    lab = ws.cell(r, 1).value
    if not lab:
        continue
    u["tco"][lab] = {"Stress": ws.cell(r, 3).value, "Base": ws.cell(r, 4).value, "Favourable": ws.cell(r, 5).value, "Ethiopia benchmark": ws.cell(r, 6).value}
res = {}
grid_rows = {}
for r in range(14, ws.max_row + 1):
    lab = ws.cell(r, 1).value
    if lab in ("NPV at the base discount rate", "IRR (pre-tax, unlevered, incl. terminal value)", "NPV with no terminal value",
               "Break-even fee at Base build cost (NPV = 0)", "Stabilised EBITDA yield on build cost (final year)",
               "NPV change per +$1/kW-month of fee", "NPV change per +$1m/MW of build cost"):
        res[lab] = ws.cell(r, 3).value
    if lab and str(lab).startswith("NPV ($m): infrastructure fee"):
        hr = r + 1
        caps = [ws.cell(hr, 3 + j).value for j in range(6)]
        rows = []
        for k in range(1, 7):
            rows.append({"fee": ws.cell(hr + k, 2).value, "npv": [ws.cell(hr + k, 3 + j).value for j in range(6)]})
        be = [ws.cell(hr + 7, 3 + j).value for j in range(6)]
        grid_rows = {"capex": caps, "rows": rows, "breakeven": be}
    if lab and str(lab).startswith("NPV ($m) at different discount rates"):
        u["wacc"] = [{"rate": ws.cell(r + k, 2).value, "npv": ws.cell(r + k, 3).value} for k in range(1, 5)]
u["campus"] = res
u["sens"] = grid_rows
# campus cash-flow rows
cf = {}
for r in range(14, ws.max_row + 1):
    lab = ws.cell(r, 1).value
    if lab in ("Calendar year", "Installed IT capacity", "Contracted IT load (capped at installed)", "Infrastructure-fee revenue", "EBITDA", "Pre-tax unlevered cash flow"):
        cf[lab] = [ws.cell(r, 3 + i).value for i in range(12)]
u["cashflow"] = cf
out["unit"] = u

# ---- Levers ----
ws = wb["Levers"]
lv = {"levers": []}
for r in range(5, 10):
    lv["levers"].append({"name": ws.cell(r, 1).value, "mw": ws.cell(r, 3).value, "lead": ws.cell(r, 4).value, "evidence": ws.cell(r, 5).value,
                         "tag": ws.cell(r, 6).value, "cost": ws.cell(r, 7).value})
lv["total_mw"] = ws.cell(10, 3).value
lv["total_cost"] = ws.cell(10, 7).value
for r in range(11, ws.max_row + 1):
    lab = ws.cell(r, 1).value
    if lab:
        lv[lab] = ws.cell(r, 3).value
out["levers"] = lv

# ---- Screen ----
ws = wb["Screen"]
crit = [ws.cell(4, c).value for c in range(2, 7)]
weights = [ws.cell(5, c).value for c in range(2, 7)]
rows = []
for r in range(6, 13):
    rows.append({"industry": ws.cell(r, 1).value, "scores": [ws.cell(r, c).value for c in range(2, 7)], "score": ws.cell(r, 7).value,
                 "rank": ws.cell(r, 8).value, "rank_equal": ws.cell(r, 9).value})
out["screen"] = {"criteria": crit, "weights": weights, "rows": rows}

# ---- Inputs (for the assumptions appendix) ----
ws = wb["Inputs"]
inputs = []
for r in range(5, ws.max_row + 1):
    pid = ws.cell(r, 1).value
    if pid and ws.cell(r, 2).value:
        inputs.append({"id": pid, "label": ws.cell(r, 2).value, "unit": ws.cell(r, 3).value,
                       "stress": ws.cell(r, 4).value, "base": ws.cell(r, 5).value, "fav": ws.cell(r, 6).value,
                       "tag": ws.cell(r, 7).value, "source": ws.cell(r, 8).value})
out["inputs"] = inputs
out["checks_ok"] = wb["Checks"].cell(wb["Checks"].max_row, 2).value

json.dump(out, open(dst, "w"), indent=1, default=str)
print("wrote", dst, "checks_ok =", out["checks_ok"])
