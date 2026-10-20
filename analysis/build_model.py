"""Build the Kenya Green Industry Hub Excel model with live formulas.

Usage:  python analysis/build_model.py [output.xlsx]
Then:   python /mnt/skills/public/xlsx/scripts/recalc.py output.xlsx   (computes cached values)

Design rules
- Every number the reader sees is either an input (blue) on the Inputs sheet or a formula.
- Inputs carry a tag: [S] sourced (with source ID), [A] analyst assumption (with rationale).
- Scenario columns: Stress / Base / Favourable for supply, economics. Hub cases use the same three
  columns relabelled Status quo / Targeted hub / Ambitious hub.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

sys.path.insert(0, __file__.rsplit("/", 1)[0])
from sources import SOURCES, CONFLICTS  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else "model/kenya_green_hub_model.xlsx"

# ---------- styles ----------
FONT = "Arial"
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True, color="12343B")
f_sub = Font(name=FONT, size=10, italic=True, color="555555")
fill_head = PatternFill("solid", fgColor="12343B")
fill_sec = PatternFill("solid", fgColor="DCE9E6")
fill_key = PatternFill("solid", fgColor="FFFF00")
fill_out = PatternFill("solid", fgColor="F3F7F6")
thin = Side(style="thin", color="BBBBBB")
b_bottom = Border(bottom=thin)

NUM = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'
PCT = '0.0%;(0.0%);-'
USD0 = '$#,##0;($#,##0);-'
USD1 = '$#,##0.0;($#,##0.0);-'
USD3 = '$0.000;($0.000);-'
YR = '0'
X1 = '0.0"x"'

wb = Workbook()

YEARS = list(range(2026, 2033))
YC = [get_column_letter(3 + i) for i in range(len(YEARS))]  # C..I
Y2030 = YC[YEARS.index(2030)]
SC = [("Stress", "D"), ("Base", "E"), ("Favourable", "F")]
HC = [("Status quo", "D"), ("Targeted hub", "E"), ("Ambitious hub", "F")]


def style_range(ws, rng, font=None, fill=None, fmt=None, align=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if fmt: c.number_format = fmt
            if align: c.alignment = align


def hdr(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = f_head
        c.fill = fill_head
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f_title
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f_sub
    ws.sheet_view.showGridLines = False


# =====================================================================
# INPUTS
# =====================================================================
wsI = wb.active
wsI.title = "Inputs"
title(wsI, "Inputs and assumptions",
      "Blue = hard-coded input. Yellow fill = key lever. Tag: [S] sourced (see Sources sheet) / [A] analyst assumption. "
      "Scenario-invariant inputs are entered once in the Base column.")
hdr(wsI, 4, ["ID", "Parameter", "Unit", "Stress", "Base", "Favourable", "Tag", "Source / rationale"])
for col, w in zip("ABCDEFGH", [14, 52, 14, 12, 12, 12, 7, 110]):
    wsI.column_dimensions[col].width = w

R = {}  # name -> row
row = 5


def section(label, cols=None):
    global row
    for ci in range(1, 9):
        wsI.cell(row=row, column=ci).fill = fill_sec
    wsI.cell(row=row, column=1, value=label).font = f_bold
    if cols:
        for lab, col in zip(cols, "DEF"):
            c = wsI[f"{col}{row}"]
            c.value = lab
            c.font = f_bold
            c.alignment = Alignment(horizontal="center")
    row += 1


def param(name, label, unit, vals, tag, src, fmt, key=False):
    """vals: single value (fixed -> Base col, D/F link to E) or tuple of 3 (stress, base, fav)."""
    global row
    R[name] = row
    wsI.cell(row=row, column=1, value=name).font = f_sub
    wsI.cell(row=row, column=2, value=label).font = f_norm
    wsI.cell(row=row, column=3, value=unit).font = f_norm
    if isinstance(vals, tuple):
        for v, col in zip(vals, "DEF"):
            c = wsI[f"{col}{row}"]
            c.value = v
            c.font = f_input
            c.number_format = fmt
            if key: c.fill = fill_key
    else:
        c = wsI[f"E{row}"]
        c.value = vals
        c.font = f_input
        c.number_format = fmt
        if key: c.fill = fill_key
        for col in "DF":
            cc = wsI[f"{col}{row}"]
            cc.value = f"=E{row}"
            cc.font = f_norm
            cc.number_format = fmt
    t = wsI.cell(row=row, column=7, value=tag)
    t.font = f_bold
    t.alignment = Alignment(horizontal="center")
    s = wsI.cell(row=row, column=8, value=src)
    s.font = f_norm
    s.alignment = Alignment(wrap_text=True, vertical="top")
    row += 1


def I(name, col="E", absolute=True):
    return f"Inputs!${col}${R[name]}" if absolute else f"Inputs!{col}{R[name]}"


section("A. Grid baseline (June 2026)", ["Stress", "Base", "Favourable"])
param("peak0", "Peak demand, June 2026", "MW", 2514, "[S]", "S2: Kenya Power peak demand 2,514 MW in June 2026 (2,549 MW in July 2026).", NUM)
param("rm0", "Reserve margin applied to the June 2026 peak", "%", (-0.015, 0.0334, 0.2073), "[S]",
      "S3: -1.5% over the 12 months to Jun 2026 (Stress). S2: 3.34% in Jun 2026 (Base). S2: 20.73% in Jan 2026 - Favourable assumes January-level availability is restored and sustained. "
      "IEA comfortable range 20-35% (S1).", PCT, key=True)
param("growth", "Peak-demand growth per year (excluding the hub programme)", "%/yr", (0.08, 0.065, 0.05), "[S/A]",
      "S2: official projection c.5%/yr (Favourable). Observed: 6.4% FY24-FY25 and 8.5% FY25-Jun 2026 (S2 data) -> Base 6.5%, Stress 8%.", PCT, key=True)

section("B. Firm, evening-capable supply additions (MW)", ["Stress", "Base", "Favourable"])
param("eth_mw", "Ethiopia import uplift (200 -> 400 MW)", "MW", (100, 200, 200), "[S/A]",
      "S2: Ethiopia imports of 200 MW 'expected to double'; S19: PPA allows up to 400 MW. Stress assumes only half the uplift is delivered.", NUM)
param("eth_yr", "Ethiopia uplift: first year counted in annual peak", "year", (2028, 2027, 2027), "[A]",
      "S2 says 'from December' without a year; Base assumes Dec 2026 so 2027 is the first full peak year (conflict C3).", YR)
param("olk1_mw", "Olkaria I rehabilitation (45 -> 61 MW)", "MW", 16, "[S]", "S13: +16 MW, late 2026.", NUM)
param("olk1_yr", "Olkaria I: first year counted", "year", (2028, 2027, 2027), "[A]", "S13 says late 2026 -> first full peak year 2027.", YR)
param("bess_mw", "Battery storage (KenGen BESS), evening-firm MW", "MW", 100, "[S]", "S13: 100 MW battery system planned for 2027; counted at full rating for a 4-hour evening peak.", NUM)
param("bess_yr", "BESS: first year counted", "year", (2028, 2027, 2027), "[S/A]", "S13: 2027. Stress delays one year.", YR)
param("bess2_mw", "Additional storage procured through auctions", "MW", (0, 0, 100), "[A]",
      "S14: BESS to be competitively sourced inside the auction framework. Only the Favourable case assumes an extra 100 MW.", NUM)
param("bess2_yr", "Additional storage: first year counted", "year", 2029, "[A]", "Earliest realistic COD given S14 timelines (auction -> PPA -> financial close -> COD).", YR)
param("olk7_mw", "Olkaria VII geothermal", "MW", 80.3, "[S]", "S13: 80.3 MW under development; commissioning date not stated.", NUM1)
param("olk7_yr", "Olkaria VII: first year counted", "year", (2030, 2029, 2028), "[A]", "S13 gives no date. Base assumes 2029.", YR)
param("solar_mw", "Seven Forks solar", "MW", 42.5, "[S]", "S13: 42.5 MW solar.", NUM1)
param("solar_firm", "Seven Forks: share counted at the evening peak", "%", 0.0, "[A]",
      "S2: solar output falls away as the 6-10pm peak arrives, so solar adds nothing to the evening reserve margin.", PCT)
param("solar_yr", "Seven Forks: first year counted", "year", 2028, "[A]", "Date not stated in S13; irrelevant while the evening contribution is 0%.", YR)
param("auction_mw", "New auction-procured firm capacity, per year from the first COD year", "MW/yr", (0, 150, 300), "[A]",
      "S14: auctions operational within 12 months; financial close 9-12 months after PPA; COD 24-30 months after financial close -> first COD c.2030. "
      "Volume is an assumption: Stress none, Base 150 MW/yr, Favourable 300 MW/yr.", NUM, key=True)
param("auction_yr", "First COD year of auction-procured capacity", "year", 2030, "[S/A]", "Derived from the S14 timeline.", YR)

section("C. Data-centre demand (hub cases)", ["Status quo", "Targeted hub", "Ambitious hub"])
param("africa2030", "Africa data-centre capacity in 2030 (midpoint of 1.5-2.2 GW)", "MW", 1850, "[S]", "S12: McKinsey range 1.5-2.2 GW by 2030; midpoint used.", NUM)
param("share", "Kenya share of Africa data-centre capacity in 2030", "%", (0.04, 0.12, 0.20), "[S/A]",
      "Status quo = Kenya's current share (c.20 MW of >500 MW, S9/S11 = 4%). Targeted = announced pipeline delivered (c.220 MW, S9). Ambitious = 20%, comparable to a clear No.2 hub.", PCT, key=True)
param("it2026", "Kenya data-centre IT capacity, 2026", "MW", 45, "[S]", "S9: almost 20 MW existing plus almost 25 MW expected by end-2025.", NUM)
param("pue", "Power usage effectiveness (PUE)", "x", 1.4, "[A]", "Typical 1.3-1.5 for a temperate-altitude site such as Nairobi; grid MW = IT MW x PUE.", '0.00')
param("util", "Average IT utilisation (flat load at the evening peak)", "%", 0.75, "[A]", "Contracted-and-drawing share of installed IT capacity at maturity; data-centre load is near-flat so peak = average.", PCT)

section("D. Economics", ["Stress", "Base", "Favourable"])
param("fee", "Colocation infrastructure fee (ex-power)", "$/kW-month", (110, 150, 200), "[A]",
      "No public Kenyan price (paywalled in S9). Range spans typical wholesale-to-retail colocation. Replace with operator quotes before relying on it.", USD0, key=True)
param("capex", "Data-centre build cost", "$m per IT MW", (11.3, 8.0, 7.0), "[S]",
      "S21: JLL global average $11.3m/MW (Stress); emerging markets below $8m/MW (Base).", USD1)
param("campus_mw", "Average campus size", "MW IT", 35, "[A]", "Nxtra Tatu City is 44 MW (S10); smaller colocation sites pull the average down.", NUM)
param("jobs", "Operational jobs per campus", "jobs", (100, 150, 200), "[S]", "S22: roughly 100-200 jobs at a typical treated US county; colocation is 'modest'.", NUM)
param("tariff", "All-in grid tariff to a large data-centre customer", "$/kWh", (0.11, 0.09, 0.065), "[S/A]",
      "S18: KSh 7.99/kWh (2020); S16: April 2026 pass-throughs add c.KSh 4.7/kWh; S15: wholesale c.$0.07. "
      "At c.KSh 129/$ (S15 implied) the range is $0.065-0.11 (conflict C5).", USD3, key=True)
param("eth_price", "Benchmark: Ethiopia-to-Kenya wholesale price", "$/kWh", 0.065, "[S]", "S19: 6.5 US cents/kWh PPA price.", USD3)
param("lcos", "Cost of shifting energy through a 4-hour battery (global benchmark)", "$/MWh", 78, "[S]", "S23: BNEF 2025 four-hour battery LCOE.", USD0)
param("em_prem", "Emerging-market premium on storage cost", "x", 1.5, "[A]", "Financing, import and logistics premium over the global benchmark.", '0.00')
param("hours", "Evening peak hours to cover", "hours/day", 4, "[S]", "S2: evening peak c.6pm-10pm.", NUM)
param("rte", "Battery round-trip efficiency", "%", 0.85, "[A]", "Typical lithium-ion utility-scale range 80-90%.", PCT)
param("charge_price", "Energy price paid to recharge the battery", "$/MWh", 70, "[S]", "S15: wholesale c.$0.07/kWh.", USD0)
param("min_rm", "Minimum reserve margin before connecting new hub load", "%", 0.10, "[A]",
      "Gate threshold. Below the IEA 20-35% comfortable range (S1) but well above the 3.3% of Jun 2026.", PCT, key=True)
param("target_rm", "Target reserve margin to restore", "%", 0.15, "[A]", "Planning target; sensitivity on the Matrix sheet.", PCT, key=True)

section("E. Reference campus (Nxtra-scale, 44 MW in two phases)", None)
param("mw1", "Phase 1 IT capacity", "MW", 22, "[S]", "S10: two phases of 22 MW.", NUM)
param("mw2", "Phase 2 IT capacity", "MW", 22, "[S]", "S10.", NUM)
param("ph2_off", "Phase 2 starts this many years after phase 1", "years", 2, "[A]", "Typical lag between phases.", NUM)
param("wacc", "Discount rate (USD, pre-tax, unlevered)", "%", 0.10, "[A]", "Hurdle for emerging-market digital infrastructure with some DFI participation. Sensitivity on UnitEconomics.", PCT, key=True)
param("opex_pct", "Variable opex (staff, maintenance, insurance, connectivity)", "% of revenue", 0.20, "[A]", "Analyst assumption.", PCT)
param("fixed_opex", "Fixed opex", "$m per installed MW per year", 0.15, "[A]", "Analyst assumption.", '0.00')
param("maint", "Maintenance capex", "% of cumulative growth capex per year", 0.015, "[A]", "Analyst assumption.", PCT)
param("tv_mult", "Terminal value at end of horizon", "x EBITDA", 7.0, "[A]", "Conservative versus 15-20 year asset life; set to 0 to see the no-terminal-value case.", X1)
param("start_year", "Campus phase 1 first operating year", "year", 2027, "[S]", "S10: ready for service Q1 2027.", YR)

section("F. 'What it would take' package (Base supply, Targeted hub)", None)
param("lv_storage", "Utility-scale 4-hour storage procured through auctions", "MW", 250, "[A]", "S14: BESS to be competitively sourced; S13: first 100 MW due 2027. Size is an assumption.", NUM)
param("lv_import", "Second Ethiopia import tranche (beyond the 400 MW PPA cap)", "MW", 150, "[A]", "S2: Ethiopia has large hydro surplus; S19: current PPA caps at 400 MW. Size is an assumption.", NUM)
param("lv_geo", "Accelerated geothermal (Olkaria VII pull-forward + mobile wellhead units)", "MW", 100, "[A]", "S13: wellhead units 'by 2027' (MW not stated); Olkaria VII 80.3 MW. Size is an assumption.", NUM)
param("lv_auction", "Auction tranche pulled forward one year", "MW", 150, "[A]", "S14: first COD c.2030 on the stated timeline; pull-forward needs procurement and financial-close discipline.", NUM)
param("lf_import", "Load factor of imported power", "%", 0.85, "[A]", "Analyst assumption for indicative energy cost.", PCT)
param("lf_geo", "Load factor of geothermal / new firm generation", "%", 0.90, "[A]", "Geothermal runs near-baseload; analyst assumption.", PCT)

wsI.freeze_panes = "C5"

# =====================================================================
# GRID
# =====================================================================
wsG = wb.create_sheet("Grid")
title(wsG, "Grid headroom, 2026-2032: firm evening-capable supply versus peak demand (before any hub load)",
      "Each block is one supply scenario. Negative reserve margin = unserved peak (rationing) unless emergency capacity is bought; it measures the size of the response needed.")
wsG.column_dimensions["A"].width = 52
wsG.column_dimensions["B"].width = 12
for c in YC:
    wsG.column_dimensions[c].width = 11

G = {}
r0 = 4
for sname, scol in SC:
    hdr(wsG, r0, [f"{sname} supply case", "Unit"] + YEARS)
    for c in YC:
        wsG[f"{c}{r0}"].number_format = YR
    rr = r0 + 1
    # peak
    G[(sname, "peak")] = rr
    wsG.cell(rr, 1, "Peak demand").font = f_bold
    wsG.cell(rr, 2, "MW").font = f_norm
    for i, c in enumerate(YC):
        if i == 0:
            wsG[f"{c}{rr}"] = f"={I('peak0', scol)}"
            wsG[f"{c}{rr}"].font = f_link
        else:
            wsG[f"{c}{rr}"] = f"={YC[i-1]}{rr}*(1+{I('growth', scol)})"
            wsG[f"{c}{rr}"].font = f_norm
        wsG[f"{c}{rr}"].number_format = NUM
    rr += 1
    # baseline firm
    G[(sname, "base")] = rr
    wsG.cell(rr, 1, "Baseline firm available capacity (Jun 2026, incl. imports)").font = f_norm
    wsG.cell(rr, 2, "MW").font = f_norm
    for c in YC:
        wsG[f"{c}{rr}"] = f"={I('peak0', scol)}*(1+{I('rm0', scol)})"
        wsG[f"{c}{rr}"].font = f_norm
        wsG[f"{c}{rr}"].number_format = NUM
    rr += 1
    adds = [
        ("Ethiopia import uplift", "eth_mw", "eth_yr", None),
        ("Olkaria I rehabilitation", "olk1_mw", "olk1_yr", None),
        ("Battery storage (KenGen BESS)", "bess_mw", "bess_yr", None),
        ("Additional storage (auctions)", "bess2_mw", "bess2_yr", None),
        ("Olkaria VII geothermal", "olk7_mw", "olk7_yr", None),
        ("Seven Forks solar (evening-firm share only)", "solar_mw", "solar_yr", "solar_firm"),
    ]
    first_add = rr
    for lab, mwn, yrn, derate in adds:
        wsG.cell(rr, 1, "  + " + lab).font = f_norm
        wsG.cell(rr, 2, "MW").font = f_norm
        for c in YC:
            core = f"{I(mwn, scol)}" + (f"*{I(derate, scol)}" if derate else "")
            wsG[f"{c}{rr}"] = f"=IF({c}${r0}>={I(yrn, scol)},{core},0)"
            wsG[f"{c}{rr}"].font = f_norm
            wsG[f"{c}{rr}"].number_format = NUM1
        rr += 1
    # auctions
    wsG.cell(rr, 1, "  + Auction-procured firm capacity (cumulative)").font = f_norm
    wsG.cell(rr, 2, "MW").font = f_norm
    for c in YC:
        wsG[f"{c}{rr}"] = f"={I('auction_mw', scol)}*MAX(0,{c}${r0}-{I('auction_yr', scol)}+1)"
        wsG[f"{c}{rr}"].font = f_norm
        wsG[f"{c}{rr}"].number_format = NUM1
    last_add = rr
    rr += 1
    # firm total
    G[(sname, "firm")] = rr
    wsG.cell(rr, 1, "Firm evening-capable capacity").font = f_bold
    wsG.cell(rr, 2, "MW").font = f_norm
    for c in YC:
        wsG[f"{c}{rr}"] = f"={c}{G[(sname,'base')]}+SUM({c}{first_add}:{c}{last_add})"
        wsG[f"{c}{rr}"].font = f_bold
        wsG[f"{c}{rr}"].number_format = NUM
    rr += 1
    # reserve margin
    G[(sname, "rm")] = rr
    wsG.cell(rr, 1, "Reserve margin before hub load").font = f_bold
    wsG.cell(rr, 2, "%").font = f_norm
    for c in YC:
        wsG[f"{c}{rr}"] = f"={c}{G[(sname,'firm')]}/{c}{G[(sname,'peak')]}-1"
        wsG[f"{c}{rr}"].font = f_bold
        wsG[f"{c}{rr}"].number_format = PCT
        wsG[f"{c}{rr}"].fill = fill_out
    rr += 1
    # headroom MW
    G[(sname, "head")] = rr
    wsG.cell(rr, 1, "Spare MW (firm capacity minus peak)").font = f_norm
    wsG.cell(rr, 2, "MW").font = f_norm
    for c in YC:
        wsG[f"{c}{rr}"] = f"={c}{G[(sname,'firm')]}-{c}{G[(sname,'peak')]}"
        wsG[f"{c}{rr}"].font = f_norm
        wsG[f"{c}{rr}"].number_format = NUM
    rr += 1
    # annual need
    G[(sname, "need")] = rr
    wsG.cell(rr, 1, "New firm MW needed each year just to hold the margin").font = f_norm
    wsG.cell(rr, 2, "MW/yr").font = f_norm
    for i, c in enumerate(YC):
        if i == 0:
            continue
        wsG[f"{c}{rr}"] = f"={c}{G[(sname,'peak')]}-{YC[i-1]}{G[(sname,'peak')]}"
        wsG[f"{c}{rr}"].font = f_norm
        wsG[f"{c}{rr}"].number_format = NUM
    rr += 2
    r0 = rr

# =====================================================================
# HUB
# =====================================================================
wsH = wb.create_sheet("Hub")
title(wsH, "Data-centre demand and the size of the prize, 2026-2032",
      "Kenya IT capacity moves linearly from 2026 to the 2030 target (share of Africa capacity) and keeps that pace afterwards. "
      "Grid load counts only capacity above the Status quo case, because Status quo growth is already inside the peak-demand growth rate.")
wsH.column_dimensions["A"].width = 56
wsH.column_dimensions["B"].width = 12
for c in YC:
    wsH.column_dimensions[c].width = 11

H = {}
r0 = 4
for hname, hcol in HC:
    hdr(wsH, r0, [f"{hname}", "Unit"] + YEARS)
    for c in YC:
        wsH[f"{c}{r0}"].number_format = YR
    rr = r0 + 1
    H[(hname, "it")] = rr
    wsH.cell(rr, 1, "Kenya data-centre IT capacity").font = f_bold
    wsH.cell(rr, 2, "MW IT").font = f_norm
    for c in YC:
        wsH[f"{c}{rr}"] = (f"={I('it2026', hcol)}+({I('share', hcol)}*{I('africa2030', hcol)}-{I('it2026', hcol)})"
                           f"*({c}${r0}-2026)/4")
        wsH[f"{c}{rr}"].font = f_bold
        wsH[f"{c}{rr}"].number_format = NUM
    rr += 1
    H[(hname, "incr")] = rr
    wsH.cell(rr, 1, "IT capacity above Status quo").font = f_norm
    wsH.cell(rr, 2, "MW IT").font = f_norm
    for c in YC:
        wsH[f"{c}{rr}"] = f"=MAX(0,{c}{H[(hname,'it')]}-{c}{4+1})"  # status quo block IT row is row 5
        wsH[f"{c}{rr}"].font = f_norm
        wsH[f"{c}{rr}"].number_format = NUM
    rr += 1
    H[(hname, "load")] = rr
    wsH.cell(rr, 1, "Hub load on the grid at the evening peak").font = f_bold
    wsH.cell(rr, 2, "MW").font = f_norm
    for c in YC:
        wsH[f"{c}{rr}"] = f"={c}{H[(hname,'incr')]}*{I('pue', hcol)}*{I('util', hcol)}"
        wsH[f"{c}{rr}"].font = f_bold
        wsH[f"{c}{rr}"].number_format = NUM
        wsH[f"{c}{rr}"].fill = fill_out
    rr += 1
    H[(hname, "rev")] = rr
    wsH.cell(rr, 1, "Infrastructure-fee revenue pool (Base fee)").font = f_norm
    wsH.cell(rr, 2, "$m/yr").font = f_norm
    for c in YC:
        wsH[f"{c}{rr}"] = f"={c}{H[(hname,'it')]}*{I('util')}*{I('fee')}*12/1000"
        wsH[f"{c}{rr}"].font = f_norm
        wsH[f"{c}{rr}"].number_format = NUM1
    rr += 1
    H[(hname, "capex")] = rr
    wsH.cell(rr, 1, "Cumulative build investment since 2026 (Base cost)").font = f_norm
    wsH.cell(rr, 2, "$m").font = f_norm
    for c in YC:
        wsH[f"{c}{rr}"] = f"=({c}{H[(hname,'it')]}-{I('it2026')})*{I('capex')}"
        wsH[f"{c}{rr}"].font = f_norm
        wsH[f"{c}{rr}"].number_format = NUM
    rr += 2
    r0 = rr

# 2030 prize summary
r = r0
hdr(wsH, r, ["2030 prize and cost, by hub case", "Unit", "Status quo", "Targeted hub", "Ambitious hub"])
r += 1
P = {}
pcols = {"Status quo": "C", "Targeted hub": "D", "Ambitious hub": "E"}


def prize_row(key, label, unit, fn, fmt):
    global r
    P[key] = r
    wsH.cell(r, 1, label).font = f_norm
    wsH.cell(r, 2, unit).font = f_norm
    for hname, hcol in HC:
        c = pcols[hname]
        wsH[f"{c}{r}"] = fn(hname, hcol)
        wsH[f"{c}{r}"].font = f_norm
        wsH[f"{c}{r}"].number_format = fmt
    r += 1


prize_row("it30", "Kenya IT capacity in 2030", "MW IT", lambda h, c: f"={Y2030}{H[(h,'it')]}", NUM)
prize_row("share30", "Share of Africa capacity in 2030", "%", lambda h, c: f"={pcols[h]}{P['it30']}/{I('africa2030')}", PCT)
prize_row("rev30", "Infrastructure-fee revenue pool, 2030", "$m/yr", lambda h, c: f"={Y2030}{H[(h,'rev')]}", NUM)
prize_row("cap30", "Build investment 2026-2030", "$bn", lambda h, c: f"={Y2030}{H[(h,'capex')]}/1000", '$0.00')
prize_row("sites", "Campuses at average size", "count", lambda h, c: f"={pcols[h]}{P['it30']}/{I('campus_mw')}", NUM1)
prize_row("jobs30", "Operational jobs (base jobs per campus)", "jobs", lambda h, c: f"={pcols[h]}{P['sites']}*{I('jobs')}", NUM)
prize_row("load30", "Hub load added to the grid, 2030", "MW", lambda h, c: f"={Y2030}{H[(h,'load')]}", NUM)
prize_row("firm_cost_mw", "Cost to firm 1 MW of evening peak with 4-hour storage", "$ per MW-yr",
          lambda h, c: f"={I('hours')}*365*({I('lcos')}*{I('em_prem')}+(1/{I('rte')}-1)*{I('charge_price')})", USD0)
prize_row("firm_cost", "Hub-attributable firming cost, 2030", "$m/yr", lambda h, c: f"={pcols[h]}{P['load30']}*{pcols[h]}{P['firm_cost_mw']}/1000000", NUM1)
prize_row("firm_ratio", "Firming cost as % of infrastructure-fee revenue", "%", lambda h, c: f"=IF({pcols[h]}{P['rev30']}=0,0,{pcols[h]}{P['firm_cost']}/{pcols[h]}{P['rev30']})", PCT)
prize_row("energy30", "Energy drawn by the whole data-centre estate, 2030", "TWh/yr",
          lambda h, c: f"={pcols[h]}{P['it30']}*{I('pue')}*{I('util')}*8760/1000000", '0.00')
prize_row("topdown", "Cross-check: projected Kenya data-centre market in 2031 (S7)", "$m", lambda h, c: "=805", NUM)
prize_row("tdratio", "Top-down market / bottom-up fee pool (2031 vs 2030)", "x", lambda h, c: f"=IF({pcols[h]}{P['rev30']}=0,0,{pcols[h]}{P['topdown']}/{pcols[h]}{P['rev30']})", X1)
wsH[f"C{P['topdown']}"].font = f_input
wsH[f"D{P['topdown']}"].font = f_input
wsH[f"E{P['topdown']}"].font = f_input
for k in ("topdown",):
    for c in "CDE":
        wsH[f"{c}{P[k]}"].comment = Comment("S7: Kenya DC market projected at $805m by 2031. Likely a broader definition than infrastructure fees (includes cloud/services).", "model")

# =====================================================================
# MATRIX
# =====================================================================
wsM = wb.create_sheet("Matrix")
title(wsM, "Supply scenario x hub ambition: reserve margin, firm-MW gap and connectable load",
      "Reserve margin = (firm capacity - peak - hub load) / (peak + hub load).")
wsM.column_dimensions["A"].width = 34
wsM.column_dimensions["B"].width = 18
for c in YC:
    wsM.column_dimensions[c].width = 11

hdr(wsM, 4, ["Supply case", "Hub case"] + YEARS)
for c in YC:
    wsM[f"{c}4"].number_format = YR
M = {}
rr = 5
for sname, scol in SC:
    for hname, hcol in HC:
        M[(sname, hname)] = rr
        wsM.cell(rr, 1, sname).font = f_norm
        wsM.cell(rr, 2, hname).font = f_norm
        for c in YC:
            wsM[f"{c}{rr}"] = (f"=(Grid!{c}{G[(sname,'firm')]}-Grid!{c}{G[(sname,'peak')]}-Hub!{c}{H[(hname,'load')]})"
                               f"/(Grid!{c}{G[(sname,'peak')]}+Hub!{c}{H[(hname,'load')]})")
            wsM[f"{c}{rr}"].font = f_norm
            wsM[f"{c}{rr}"].number_format = PCT
        rr += 1
rr += 1


def matrix3(label, fn, fmt, key):
    """3x3 table (rows = supply case, cols = hub case) for 2030."""
    global rr
    hdr(wsM, rr, [label, "Supply case \\ Hub case", "Status quo", "Targeted hub", "Ambitious hub"])
    rr += 1
    for sname, scol in SC:
        wsM.cell(rr, 1, "").font = f_norm
        wsM.cell(rr, 2, sname).font = f_bold
        for j, (hname, hcol) in enumerate(HC):
            c = get_column_letter(3 + j)
            wsM[f"{c}{rr}"] = fn(sname, hname, scol)
            wsM[f"{c}{rr}"].font = f_norm
            wsM[f"{c}{rr}"].number_format = fmt
            wsM[f"{c}{rr}"].fill = fill_out
            M[(key, sname, hname)] = f"{c}{rr}"
        rr += 1
    rr += 1


matrix3("Reserve margin in 2030", lambda s, h, sc: f"={Y2030}{M[(s,h)]}", PCT, "rm2030")
matrix3("Firm MW needed in 2030 to reach the minimum reserve margin",
        lambda s, h, sc: (f"=MAX(0,(1+{I('min_rm')})*(Grid!{Y2030}{G[(s,'peak')]}+Hub!{Y2030}{H[(h,'load')]})-Grid!{Y2030}{G[(s,'firm')]})"),
        NUM, "gapmin")
matrix3("Firm MW needed in 2030 to reach the target reserve margin",
        lambda s, h, sc: (f"=MAX(0,(1+{I('target_rm')})*(Grid!{Y2030}{G[(s,'peak')]}+Hub!{Y2030}{H[(h,'load')]})-Grid!{Y2030}{G[(s,'firm')]})"),
        NUM, "gaptgt")
matrix3("Storage (MWh, 4-hour) to carry the hub's own evening load in 2030",
        lambda s, h, sc: f"=Hub!{Y2030}{H[(h,'load')]}*{I('hours')}", NUM, "mwh")

# connectable load
hdr(wsM, rr, ["Connectable new data-centre IT capacity (MW) while keeping the minimum reserve margin", "Supply case"] + YEARS)
for c in YC:
    wsM[f"{c}{rr}"].number_format = YR
rr += 1
CN = {}
for sname, scol in SC:
    CN[sname] = rr
    wsM.cell(rr, 2, sname).font = f_bold
    for c in YC:
        wsM[f"{c}{rr}"] = (f"=MAX(0,(Grid!{c}{G[(sname,'firm')]}/(1+{I('min_rm')})-Grid!{c}{G[(sname,'peak')]})"
                           f"/({I('pue')}*{I('util')}))")
        wsM[f"{c}{rr}"].font = f_norm
        wsM[f"{c}{rr}"].number_format = NUM
        wsM[f"{c}{rr}"].fill = fill_out
    rr += 1
rr += 1
hdr(wsM, rr, ["Connectable new data-centre IT capacity (MW) at the target reserve margin", "Supply case"] + YEARS)
for c in YC:
    wsM[f"{c}{rr}"].number_format = YR
rr += 1
CT = {}
for sname, scol in SC:
    CT[sname] = rr
    wsM.cell(rr, 2, sname).font = f_bold
    for c in YC:
        wsM[f"{c}{rr}"] = (f"=MAX(0,(Grid!{c}{G[(sname,'firm')]}/(1+{I('target_rm')})-Grid!{c}{G[(sname,'peak')]})"
                           f"/({I('pue')}*{I('util')}))")
        wsM[f"{c}{rr}"].font = f_norm
        wsM[f"{c}{rr}"].number_format = NUM
    rr += 1

# =====================================================================
# UNIT ECONOMICS
# =====================================================================
wsU = wb.create_sheet("UnitEconomics")
title(wsU, "Unit economics: what a tenant pays, and whether a reference campus clears its hurdle",
      "Part 1 is a tenant's all-in cost per IT kW-month. Part 2 is an illustrative pre-tax, unlevered campus case; every driver is an input.")
wsU.column_dimensions["A"].width = 58
wsU.column_dimensions["B"].width = 14
for i in range(3, 16):
    wsU.column_dimensions[get_column_letter(i)].width = 11

U = {}
hdr(wsU, 4, ["Part 1. Tenant cost per IT kW-month", "Unit", "Stress", "Base", "Favourable", "Ethiopia-linked benchmark"])
rows1 = [
    ("fee", "Infrastructure fee", "$/kW-month", lambda col: f"={I('fee', col)}", lambda: f"={I('fee','E')}"),
    ("power", "Grid power (PUE x 730 h x all-in tariff)", "$/kW-month", lambda col: f"={I('pue', col)}*730*{I('tariff', col)}", lambda: f"={I('pue','E')}*730*{I('eth_price','E')}"),
    ("evening", "Evening self-supply premium (storage shifting cost, per IT kW)", "$/kW-month",
     lambda col: f"={I('pue', col)}*{I('hours', col)}*365/12*({I('lcos', col)}*{I('em_prem', col)}+(1/{I('rte', col)}-1)*{I('charge_price', col)})/1000",
     lambda: f"={I('pue','E')}*{I('hours','E')}*365/12*({I('lcos','E')}*{I('em_prem','E')}+(1/{I('rte','E')}-1)*{I('charge_price','E')})/1000"),
]
rr = 5
for key, lab, unit, fn, bench in rows1:
    U[key] = rr
    wsU.cell(rr, 1, lab).font = f_norm
    wsU.cell(rr, 2, unit).font = f_norm
    for (sname, scol), c in zip(SC, "CDE"):
        wsU[f"{c}{rr}"] = fn(scol)
        wsU[f"{c}{rr}"].number_format = USD1
        wsU[f"{c}{rr}"].font = f_norm
    wsU[f"F{rr}"] = bench()
    wsU[f"F{rr}"].number_format = USD1
    wsU[f"F{rr}"].font = f_norm
    rr += 1
U["tco"] = rr
wsU.cell(rr, 1, "All-in tenant cost").font = f_bold
wsU.cell(rr, 2, "$/kW-month").font = f_norm
for c in "CDEF":
    wsU[f"{c}{rr}"] = f"=SUM({c}{U['fee']}:{c}{U['evening']})"
    wsU[f"{c}{rr}"].number_format = USD1
    wsU[f"{c}{rr}"].font = f_bold
    wsU[f"{c}{rr}"].fill = fill_out
rr += 1
U["pshare"] = rr
wsU.cell(rr, 1, "Power (grid + evening premium) as % of all-in cost").font = f_norm
wsU.cell(rr, 2, "%").font = f_norm
for c in "CDEF":
    wsU[f"{c}{rr}"] = f"=({c}{U['power']}+{c}{U['evening']})/{c}{U['tco']}"
    wsU[f"{c}{rr}"].number_format = PCT
    wsU[f"{c}{rr}"].font = f_norm
rr += 1
U["gap"] = rr
wsU.cell(rr, 1, "Power-price gap to the Ethiopia-linked benchmark (grid power only)").font = f_norm
wsU.cell(rr, 2, "$/kW-month").font = f_norm
for c in "CDE":
    wsU[f"{c}{rr}"] = f"={c}{U['power']}-$F${U['power']}"
    wsU[f"{c}{rr}"].number_format = USD1
    wsU[f"{c}{rr}"].font = f_norm
rr += 1
U["gappct"] = rr
wsU.cell(rr, 1, "  ...as % of the tenant's all-in cost").font = f_norm
wsU.cell(rr, 2, "%").font = f_norm
for c in "CDE":
    wsU[f"{c}{rr}"] = f"={c}{U['gap']}/{c}{U['tco']}"
    wsU[f"{c}{rr}"].number_format = PCT
    wsU[f"{c}{rr}"].font = f_norm
rr += 2

# Part 2: campus cash flow (Base case)
NT = 12
CC = [get_column_letter(3 + i) for i in range(NT)]  # C..N
hdr(wsU, rr, ["Part 2. Reference campus cash flow (Base case, $m unless stated)", "Unit"] + [f"t={i}" for i in range(NT)])
rr += 1
U["t"] = rr
wsU.cell(rr, 1, "Period index t").font = f_norm
for i, c in enumerate(CC):
    wsU[f"{c}{rr}"] = i
    wsU[f"{c}{rr}"].font = f_input
    wsU[f"{c}{rr}"].number_format = '0'
rr += 1
U["year"] = rr
wsU.cell(rr, 1, "Calendar year").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"={I('start_year')}+{c}{U['t']}"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = YR
rr += 1
U["inst"] = rr
wsU.cell(rr, 1, "Installed IT capacity").font = f_norm
wsU.cell(rr, 2, "MW").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"={I('mw1')}+IF({c}{U['t']}>={I('ph2_off')},{I('mw2')},0)"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = NUM
rr += 1
U["sold_in"] = rr
wsU.cell(rr, 1, "Contracted IT load (input ramp)").font = f_norm
wsU.cell(rr, 2, "MW").font = f_norm
ramp = [8, 15, 22, 28, 33, 33, 33, 33, 33, 33, 33, 33]
for c, v in zip(CC, ramp):
    wsU[f"{c}{rr}"] = v
    wsU[f"{c}{rr}"].font = f_input
    wsU[f"{c}{rr}"].number_format = NUM
wsU[f"A{rr}"].comment = Comment("[A] Analyst assumption: ramp to c.75% of 44 MW by year 5 (matches Inputs 'util').", "model")
rr += 1
U["sold"] = rr
wsU.cell(rr, 1, "Contracted IT load (capped at installed)").font = f_norm
wsU.cell(rr, 2, "MW").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"=MIN({c}{U['sold_in']},{c}{U['inst']})"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = NUM
rr += 1
U["rev"] = rr
wsU.cell(rr, 1, "Infrastructure-fee revenue").font = f_norm
wsU.cell(rr, 2, "$m").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"={c}{U['sold']}*{I('fee')}*12/1000"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = NUM1
rr += 1
U["vopex"] = rr
wsU.cell(rr, 1, "Variable opex").font = f_norm
wsU.cell(rr, 2, "$m").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"=-{c}{U['rev']}*{I('opex_pct')}"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = NUM1
rr += 1
U["fopex"] = rr
wsU.cell(rr, 1, "Fixed opex").font = f_norm
wsU.cell(rr, 2, "$m").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"=-{c}{U['inst']}*{I('fixed_opex')}"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = NUM1
rr += 1
U["ebitda"] = rr
wsU.cell(rr, 1, "EBITDA").font = f_bold
wsU.cell(rr, 2, "$m").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"={c}{U['rev']}+{c}{U['vopex']}+{c}{U['fopex']}"
    wsU[f"{c}{rr}"].font = f_bold
    wsU[f"{c}{rr}"].number_format = NUM1
rr += 1
U["gmw"] = rr
wsU.cell(rr, 1, "Growth capacity added (MW)").font = f_norm
wsU.cell(rr, 2, "MW").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"=IF({c}{U['t']}=0,{I('mw1')},0)+IF({c}{U['t']}={I('ph2_off')},{I('mw2')},0)"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = NUM
rr += 1
U["gcap"] = rr
wsU.cell(rr, 1, "Growth capex").font = f_norm
wsU.cell(rr, 2, "$m").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"=-{c}{U['gmw']}*{I('capex')}"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = NUM1
rr += 1
U["cummw"] = rr
wsU.cell(rr, 1, "Cumulative growth MW at start of period").font = f_norm
wsU.cell(rr, 2, "MW").font = f_norm
for i, c in enumerate(CC):
    wsU[f"{c}{rr}"] = "=0" if i == 0 else f"=SUM($C{U['gmw']}:{CC[i-1]}{U['gmw']})"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = NUM
rr += 1
U["mcap"] = rr
wsU.cell(rr, 1, "Maintenance capex").font = f_norm
wsU.cell(rr, 2, "$m").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"=-{c}{U['cummw']}*{I('capex')}*{I('maint')}"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = NUM1
rr += 1
U["tv"] = rr
wsU.cell(rr, 1, "Terminal value (multiple of final-year EBITDA)").font = f_norm
wsU.cell(rr, 2, "$m").font = f_norm
for i, c in enumerate(CC):
    wsU[f"{c}{rr}"] = f"={I('tv_mult')}*{c}{U['ebitda']}" if i == NT - 1 else "=0"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = NUM1
rr += 1
U["cf"] = rr
wsU.cell(rr, 1, "Pre-tax unlevered cash flow").font = f_bold
wsU.cell(rr, 2, "$m").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"={c}{U['ebitda']}+{c}{U['gcap']}+{c}{U['mcap']}+{c}{U['tv']}"
    wsU[f"{c}{rr}"].font = f_bold
    wsU[f"{c}{rr}"].number_format = NUM1
    wsU[f"{c}{rr}"].fill = fill_out
rr += 2

# helper rows for closed-form linearity
hdr(wsU, rr, ["Linear decomposition (NPV is linear in fee and in capex per MW)", "Unit"] + [f"t={i}" for i in range(NT)])
rr += 1
U["h_fee"] = rr
wsU.cell(rr, 1, "Fee-linked cash flow per $1/kW-month of fee").font = f_norm
wsU.cell(rr, 2, "$m").font = f_norm
for i, c in enumerate(CC):
    tv = f"+{I('tv_mult')}*{c}{U['sold']}*12/1000*(1-{I('opex_pct')})" if i == NT - 1 else ""
    wsU[f"{c}{rr}"] = f"={c}{U['sold']}*12/1000*(1-{I('opex_pct')}){tv}"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = '0.000'
rr += 1
U["h_fix"] = rr
wsU.cell(rr, 1, "Fixed-opex cash flow (incl. terminal value effect)").font = f_norm
wsU.cell(rr, 2, "$m").font = f_norm
for i, c in enumerate(CC):
    tv = f"-{I('tv_mult')}*{c}{U['inst']}*{I('fixed_opex')}" if i == NT - 1 else ""
    wsU[f"{c}{rr}"] = f"=-{c}{U['inst']}*{I('fixed_opex')}{tv}"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = '0.000'
rr += 1
U["h_cap"] = rr
wsU.cell(rr, 1, "Capex cash flow per $1m of build cost per MW").font = f_norm
wsU.cell(rr, 2, "$m").font = f_norm
for c in CC:
    wsU[f"{c}{rr}"] = f"=-({c}{U['gmw']}+{c}{U['cummw']}*{I('maint')})"
    wsU[f"{c}{rr}"].font = f_norm
    wsU[f"{c}{rr}"].number_format = '0.000'
rr += 2

hdr(wsU, rr, ["Results (Base case)", "Unit", "Value"])
rr += 1


def result(key, label, unit, formula, fmt, bold=False):
    global rr
    U[key] = rr
    wsU.cell(rr, 1, label).font = f_bold if bold else f_norm
    wsU.cell(rr, 2, unit).font = f_norm
    wsU[f"C{rr}"] = formula
    wsU[f"C{rr}"].number_format = fmt
    wsU[f"C{rr}"].font = f_bold if bold else f_norm
    wsU[f"C{rr}"].fill = fill_out
    rr += 1


cfrow = f"C{U['cf']}:N{U['cf']}"
result("npv", "NPV at the base discount rate", "$m", f"=C{U['cf']}+NPV({I('wacc')},D{U['cf']}:N{U['cf']})", NUM1, True)
result("irr", "IRR (pre-tax, unlevered, incl. terminal value)", "%", f"=IRR({cfrow})", PCT, True)
result("npv_notv", "NPV with no terminal value", "$m", f"=C{U['npv']}-N{U['tv']}/(1+{I('wacc')})^11", NUM1)
result("slope", "NPV change per +$1/kW-month of fee", "$m", f"=C{U['h_fee']}+NPV({I('wacc')},D{U['h_fee']}:N{U['h_fee']})", '0.000')
result("c0", "NPV excluding fee and capex effects (fixed opex)", "$m", f"=C{U['h_fix']}+NPV({I('wacc')},D{U['h_fix']}:N{U['h_fix']})", NUM1)
result("capcoef", "NPV change per +$1m/MW of build cost", "$m", f"=C{U['h_cap']}+NPV({I('wacc')},D{U['h_cap']}:N{U['h_cap']})", '0.000')
result("be_fee", "Break-even fee at Base build cost (NPV = 0)", "$/kW-month", f"=-(C{U['c0']}+C{U['capcoef']}*{I('capex')})/C{U['slope']}", USD1, True)
result("chk", "Check: linear decomposition reproduces NPV (should be 0)", "$m", f"=ROUND(C{U['slope']}*{I('fee')}+C{U['c0']}+C{U['capcoef']}*{I('capex')}-C{U['npv']},6)", '0.000000')
result("yield", "Stabilised EBITDA yield on build cost (final year)", "%", f"=N{U['ebitda']}/({I('capex')}*({I('mw1')}+{I('mw2')}))", PCT)
rr += 1

# sensitivity grid
hdr(wsU, rr, ["NPV ($m): infrastructure fee ($/kW-month, down) x build cost ($m per MW, across)", "Fee \\ Build cost"] + [None] * 6)
caps = [6.0, 7.0, 8.0, 9.0, 10.0, 11.3]
fees = [110, 130, 150, 175, 200, 225]
rr += 1
U["grid_hdr"] = rr
for j, cv in enumerate(caps):
    c = get_column_letter(3 + j)
    wsU[f"{c}{rr}"] = cv
    wsU[f"{c}{rr}"].font = f_input
    wsU[f"{c}{rr}"].number_format = USD1
wsU.cell(rr, 2, "").font = f_norm
rr += 1
U["grid0"] = rr
for fv in fees:
    wsU.cell(rr, 2, fv).font = f_input
    wsU.cell(rr, 2).number_format = USD0
    for j in range(len(caps)):
        c = get_column_letter(3 + j)
        wsU[f"{c}{rr}"] = f"=$C${U['slope']}*$B{rr}+$C${U['c0']}+$C${U['capcoef']}*{c}${U['grid_hdr']}"
        wsU[f"{c}{rr}"].font = f_norm
        wsU[f"{c}{rr}"].number_format = NUM
    rr += 1
wsU.cell(rr, 2, "Break-even fee").font = f_bold
for j in range(len(caps)):
    c = get_column_letter(3 + j)
    wsU[f"{c}{rr}"] = f"=-($C${U['c0']}+$C${U['capcoef']}*{c}${U['grid_hdr']})/$C${U['slope']}"
    wsU[f"{c}{rr}"].font = f_bold
    wsU[f"{c}{rr}"].number_format = USD0
    wsU[f"{c}{rr}"].fill = fill_out
U["grid_be"] = rr
rr += 2

hdr(wsU, rr, ["NPV ($m) at different discount rates (Base fee and build cost)", "Rate", "NPV"])
rr += 1
U["wacc0"] = rr
for w in [0.08, 0.10, 0.12, 0.14]:
    wsU.cell(rr, 2, w).font = f_input
    wsU.cell(rr, 2).number_format = PCT
    wsU[f"C{rr}"] = f"=$C${U['cf']}+NPV(B{rr},$D${U['cf']}:$N${U['cf']})"
    wsU[f"C{rr}"].font = f_norm
    wsU[f"C{rr}"].number_format = NUM1
    rr += 1

# =====================================================================
# LEVERS
# =====================================================================
wsL = wb.create_sheet("Levers")
title(wsL, "What it would take: five levers sized against the 2030 firm-capacity gap (Base supply, Targeted hub)",
      "Costs are indicative annual procurement costs, not a budget. Lever sizes are inputs on the Inputs sheet (section F).")
for col, w in zip("ABCDEFG", [58, 12, 14, 24, 58, 8, 16]):
    wsL.column_dimensions[col].width = w
hdr(wsL, 4, ["Lever", "Unit", "Firm MW", "Lead time", "Evidence / basis", "Tag", "Indicative cost ($m/yr)"])
L = {}
levers = [
    ("lv_hub", "1. Data centres bring their own evening power (on-site storage)", f"=Hub!D{P['load30']}", "Built with each campus (12-18 months)",
     "S2: evening peak 6-10pm; sized at the Targeted-hub load added in 2030. Cost borne by tenants.", "[A]",
     f"=C{{r}}*Hub!D{P['firm_cost_mw']}/1000000"),
    ("lv_stor", "2. Utility-scale 4-hour storage via auctions", f"={I('lv_storage')}", "24-30 months after financial close (S14)",
     "S13 first 100 MW in 2027; S14 BESS inside the auction framework.", "[A]", f"=C{{r}}*Hub!D{P['firm_cost_mw']}/1000000"),
    ("lv_imp", "3. Second Ethiopia import tranche", f"={I('lv_import')}", "12-24 months (existing interconnector)",
     "S2 / S19: 200 MW now, 400 MW PPA cap, large Ethiopian hydro surplus. Priced at the 6.5 c/kWh PPA benchmark.", "[A]",
     f"=C{{r}}*8.76*{I('lf_import')}*{I('eth_price')}"),
    ("lv_geo", "4. Accelerated geothermal (Olkaria VII, wellhead units)", f"={I('lv_geo')}", "12-36 months (S13: wellhead units by 2027)",
     "S13: Olkaria VII 80.3 MW; mobile wellhead units. Priced at wholesale $0.07/kWh (S15).", "[A]",
     f"=C{{r}}*8.76*{I('lf_geo')}*{I('charge_price')}/1000"),
    ("lv_auc", "5. Auction tranche pulled forward one year", f"={I('lv_auction')}", "Needs auction live by late 2026 (S14)",
     "S14: first COD c.2030; this lever brings one tranche to 2029. Priced at wholesale $0.07/kWh.", "[A]",
     f"=C{{r}}*8.76*{I('lf_geo')}*{I('charge_price')}/1000"),
]
rr = 5
for key, lab, mwf, lead, ev, tag, costf in levers:
    L[key] = rr
    wsL.cell(rr, 1, lab).font = f_norm
    wsL.cell(rr, 2, "MW").font = f_norm
    wsL[f"C{rr}"] = mwf
    wsL[f"C{rr}"].font = f_link
    wsL[f"C{rr}"].number_format = NUM
    wsL.cell(rr, 4, lead).font = f_norm
    wsL.cell(rr, 5, ev).font = f_norm
    wsL.cell(rr, 6, tag).font = f_bold
    wsL[f"G{rr}"] = costf.replace("{r}", str(rr))
    wsL[f"G{rr}"].font = f_norm
    wsL[f"G{rr}"].number_format = NUM1
    for ci in (1, 4, 5):
        wsL.cell(rr, ci).alignment = Alignment(wrap_text=True, vertical="top")
    rr += 1
L["total"] = rr
wsL.cell(rr, 1, "Package total").font = f_bold
wsL[f"C{rr}"] = f"=SUM(C5:C{rr-1})"
wsL[f"G{rr}"] = f"=SUM(G5:G{rr-1})"
for c, fmt in (("C", NUM), ("G", NUM1)):
    wsL[f"{c}{rr}"].font = f_bold
    wsL[f"{c}{rr}"].number_format = fmt
    wsL[f"{c}{rr}"].fill = fill_out
rr += 2
hdr(wsL, rr, ["Coverage of the 2030 gap", "Unit", "Value"])
rr += 1
L["gapmin"] = rr
wsL.cell(rr, 1, "Gap to the minimum reserve margin (Base supply, Targeted hub, 2030)").font = f_norm
wsL.cell(rr, 2, "MW").font = f_norm
wsL[f"C{rr}"] = f"=Matrix!{M[('gapmin','Base','Targeted hub')]}"
wsL[f"C{rr}"].font = f_link
wsL[f"C{rr}"].number_format = NUM
rr += 1
L["gaptgt"] = rr
wsL.cell(rr, 1, "Gap to the target reserve margin (Base supply, Targeted hub, 2030)").font = f_norm
wsL.cell(rr, 2, "MW").font = f_norm
wsL[f"C{rr}"] = f"=Matrix!{M[('gaptgt','Base','Targeted hub')]}"
wsL[f"C{rr}"].font = f_link
wsL[f"C{rr}"].number_format = NUM
rr += 1
L["cov_min"] = rr
wsL.cell(rr, 1, "Package coverage of the minimum-margin gap").font = f_bold
wsL.cell(rr, 2, "x").font = f_norm
wsL[f"C{rr}"] = f"=C{L['total']}/C{L['gapmin']}"
wsL[f"C{rr}"].number_format = '0.00"x"'
wsL[f"C{rr}"].font = f_bold
rr += 1
L["cov_tgt"] = rr
wsL.cell(rr, 1, "Package coverage of the target-margin gap").font = f_bold
wsL.cell(rr, 2, "x").font = f_norm
wsL[f"C{rr}"] = f"=C{L['total']}/C{L['gaptgt']}"
wsL[f"C{rr}"].number_format = '0.00"x"'
wsL[f"C{rr}"].font = f_bold
rr += 1
L["rm_after"] = rr
wsL.cell(rr, 1, "Reserve margin in 2030 after the package (Base supply, Targeted hub)").font = f_bold
wsL.cell(rr, 2, "%").font = f_norm
wsL[f"C{rr}"] = (f"=(Grid!{Y2030}{G[('Base','firm')]}+C{L['total']}-Grid!{Y2030}{G[('Base','peak')]}-Hub!{Y2030}{H[('Targeted hub','load')]})"
                 f"/(Grid!{Y2030}{G[('Base','peak')]}+Hub!{Y2030}{H[('Targeted hub','load')]})")
wsL[f"C{rr}"].number_format = PCT
wsL[f"C{rr}"].font = f_bold
rr += 1
L["self_share"] = rr
wsL.cell(rr, 1, "Share of the package supplied by the data centres themselves (lever 1)").font = f_norm
wsL.cell(rr, 2, "%").font = f_norm
wsL[f"C{rr}"] = f"=C{L['lv_hub']}/C{L['total']}"
wsL[f"C{rr}"].number_format = PCT
wsL[f"C{rr}"].font = f_norm
rr += 1
L["cost_pct"] = rr
wsL.cell(rr, 1, "Share of the target-margin gap attributable to the hub (hub load / gap)").font = f_norm
wsL.cell(rr, 2, "%").font = f_norm
wsL[f"C{rr}"] = f"=Hub!D{P['load30']}/C{L['gaptgt']}"
wsL[f"C{rr}"].number_format = PCT
wsL[f"C{rr}"].font = f_norm

# =====================================================================
# SCREEN
# =====================================================================
wsX = wb.create_sheet("Screen")
title(wsX, "Beyond data centres: which energy-intensive industries fit Kenya's grid shape?",
      "Scores 1-5 are analyst JUDGEMENT [J], not data. Change the blue weights and scores to test the ranking; the equal-weight ranking shows how much the weights matter.")
wsX.column_dimensions["A"].width = 52
for c in "BCDEFGHI":
    wsX.column_dimensions[c].width = 15
crit = ["Fit with an evening-constrained grid (load flexibility)", "Market pull evidence", "Kenya's input advantage", "Jobs per MW", "Speed to deploy"]
hdr(wsX, 4, ["Industry"] + crit + ["Weighted score", "Rank", "Equal-weight rank"])
wsX.row_dimensions[4].height = 54
wsX.cell(5, 1, "Weight").font = f_bold
weights = [0.30, 0.20, 0.20, 0.15, 0.15]
for j, w in enumerate(weights):
    c = get_column_letter(2 + j)
    wsX[f"{c}5"] = w
    wsX[f"{c}5"].font = f_input
    wsX[f"{c}5"].number_format = '0%'
wsX["G5"] = "=SUM(B5:F5)"
wsX["G5"].number_format = '0%'
wsX["G5"].font = f_bold
inds = [
    ("Agro-processing and cold chain", [4, 4, 4, 5, 5]),
    ("Night-shift heavy industry (cement, steel recycling, glass)", [5, 4, 3, 4, 4]),
    ("Textiles and apparel in export zones", [3, 3, 3, 5, 4]),
    ("Flexible / batch compute (AI training, rendering)", [5, 3, 3, 1, 3]),
    ("EV and battery assembly, smart charging", [4, 3, 3, 3, 3]),
    ("Cloud and colocation data centres", [2, 5, 4, 2, 3]),
    ("Green hydrogen / ammonia (electrolysis)", [5, 2, 3, 2, 1]),
]
X = {}
r1 = 6
for k, (name, sc) in enumerate(inds):
    rr = r1 + k
    X[name] = rr
    wsX.cell(rr, 1, name).font = f_norm
    for j, v in enumerate(sc):
        c = get_column_letter(2 + j)
        wsX[f"{c}{rr}"] = v
        wsX[f"{c}{rr}"].font = f_input
    wsX[f"G{rr}"] = f"=SUMPRODUCT($B$5:$F$5,B{rr}:F{rr})"
    wsX[f"G{rr}"].number_format = '0.00'
    wsX[f"G{rr}"].font = f_bold
    wsX[f"G{rr}"].fill = fill_out
rN = r1 + len(inds) - 1
for k in range(len(inds)):
    rr = r1 + k
    wsX[f"H{rr}"] = f"=RANK(G{rr},$G${r1}:$G${rN})"
    wsX[f"H{rr}"].font = f_norm
    wsX[f"I{rr}"] = f"=RANK(AVERAGE(B{rr}:F{rr}),$J${r1}:$J${rN})"
    wsX[f"J{rr}"] = f"=AVERAGE(B{rr}:F{rr})"
    wsX[f"J{rr}"].number_format = '0.00'
    wsX[f"J{rr}"].font = f_norm
    wsX[f"I{rr}"].font = f_norm
wsX["J4"] = "Equal-weight mean"
wsX["J4"].font = f_head
wsX["J4"].fill = fill_head
wsX["J4"].alignment = Alignment(horizontal="center", wrap_text=True)
wsX.column_dimensions["J"].width = 15
wsX.cell(rN + 2, 1, "Why these criteria: Kenya's binding constraint is firm power at 6-10pm (S2). Loads that can move off that window, or store their own evening energy, use the grid's idle hours "
                    "(night tariff uptake +74% in H2 2025, S17) instead of competing for its scarcest ones.").font = f_sub

# =====================================================================
# SUMMARY (front sheet)
# =====================================================================
wsS = wb.create_sheet("Summary", 0)
title(wsS, "Summary of results - Kenya as East Africa's low-carbon industry hub",
      "All figures are live links to the calculation sheets. Scenario definitions are on the Inputs sheet.")
wsS.column_dimensions["A"].width = 70
wsS.column_dimensions["B"].width = 16
wsS.column_dimensions["C"].width = 16
wsS.column_dimensions["D"].width = 16
wsS.column_dimensions["E"].width = 60
S = {}
rr = 4


def srow(key, label, vals, fmt, note=""):
    """vals: list of up to 3 formulas -> columns B,C,D"""
    global rr
    S[key] = rr
    wsS.cell(rr, 1, label).font = f_norm
    for v, c in zip(vals, "BCD"):
        wsS[f"{c}{rr}"] = v
        wsS[f"{c}{rr}"].number_format = fmt
        wsS[f"{c}{rr}"].font = f_link
    wsS.cell(rr, 5, note).font = f_sub
    rr += 1


hdr(wsS, rr, ["1. Grid headroom (supply cases)", "Stress", "Base", "Favourable"])
rr += 1
srow("peak26", "Peak demand, Jun 2026 (MW)", [f"=Grid!C{G[(s,'peak')]}" for s, _ in SC], NUM)
srow("firm26", "Firm available capacity, Jun 2026 (MW)", [f"=Grid!C{G[(s,'firm')]}" for s, _ in SC], NUM)
srow("head26", "Spare MW, Jun 2026", [f"=Grid!C{G[(s,'head')]}" for s, _ in SC], NUM, "Base: the whole national spare margin")
srow("rm27", "Reserve margin 2027 (no hub)", [f"=Grid!D{G[(s,'rm')]}" for s, _ in SC], PCT)
srow("rm28", "Reserve margin 2028 (no hub)", [f"=Grid!E{G[(s,'rm')]}" for s, _ in SC], PCT)
srow("rm30", "Reserve margin 2030 (no hub)", [f"=Grid!{Y2030}{G[(s,'rm')]}" for s, _ in SC], PCT)
srow("rm32", "Reserve margin 2032 (no hub)", [f"=Grid!I{G[(s,'rm')]}" for s, _ in SC], PCT)
srow("need30", "New firm MW needed per year just to hold the margin, 2030", [f"=Grid!{Y2030}{G[(s,'need')]}" for s, _ in SC], NUM)
rr += 1
hdr(wsS, rr, ["2. The prize (hub cases, 2030)", "Status quo", "Targeted hub", "Ambitious hub"])
rr += 1
for key, label, fmt in [("it30", "Kenya data-centre IT capacity (MW)", NUM), ("share30", "Share of Africa capacity", PCT),
                        ("rev30", "Infrastructure-fee revenue pool ($m/yr)", NUM), ("cap30", "Build investment 2026-2030 ($bn)", '$0.00'),
                        ("jobs30", "Operational jobs", NUM), ("load30", "Hub load added to grid (MW)", NUM),
                        ("firm_cost", "Hub-attributable firming cost ($m/yr)", NUM1), ("firm_ratio", "Firming cost as % of fee revenue", PCT),
                        ("energy30", "Energy drawn by the data-centre estate (TWh/yr)", '0.00')]:
    srow(key, label, [f"=Hub!{c}{P[key]}" for c in "CDE"], fmt)
rr += 1
hdr(wsS, rr, ["3. Reserve margin in 2030 (rows: supply case; columns: hub case)", "Status quo", "Targeted hub", "Ambitious hub"])
rr += 1
for sname, _ in SC:
    srow(f"rm30_{sname}", f"{sname} supply", [f"=Matrix!{M[('rm2030', sname, h)]}" for h, _ in HC], PCT)
rr += 1
hdr(wsS, rr, ["4. Firm MW needed in 2030 to reach the minimum reserve margin", "Status quo", "Targeted hub", "Ambitious hub"])
rr += 1
for sname, _ in SC:
    srow(f"gapmin_{sname}", f"{sname} supply", [f"=Matrix!{M[('gapmin', sname, h)]}" for h, _ in HC], NUM)
rr += 1
hdr(wsS, rr, ["5. Firm MW needed in 2030 to reach the target reserve margin", "Status quo", "Targeted hub", "Ambitious hub"])
rr += 1
for sname, _ in SC:
    srow(f"gaptgt_{sname}", f"{sname} supply", [f"=Matrix!{M[('gaptgt', sname, h)]}" for h, _ in HC], NUM)
rr += 1
hdr(wsS, rr, ["6. Connectable IT MW at the minimum reserve margin", "2027", "2028", "2030"])
rr += 1
for sname, _ in SC:
    srow(f"conn_{sname}", f"{sname} supply", [f"=Matrix!D{CN[sname]}", f"=Matrix!E{CN[sname]}", f"=Matrix!{Y2030}{CN[sname]}"], NUM)
rr += 1
hdr(wsS, rr, ["7. Tenant cost per IT kW-month ($)", "Stress", "Base", "Favourable"])
rr += 1
for key, label in [("fee", "Infrastructure fee"), ("power", "Grid power"), ("evening", "Evening self-supply premium"), ("tco", "All-in cost"),
                   ("pshare", "Power share of all-in cost"), ("gap", "Power-price gap vs Ethiopia-linked benchmark"), ("gappct", "  as % of all-in cost")]:
    fmt = PCT if key in ("pshare", "gappct") else USD1
    srow("u_" + key, label, [f"=UnitEconomics!{c}{U[key]}" for c in "CDE"], fmt)
rr += 1
hdr(wsS, rr, ["8. Reference campus (44 MW, Base case)", "Value", "", ""])
rr += 1
srow("npv", "NPV ($m)", [f"=UnitEconomics!C{U['npv']}"], NUM1)
srow("irr", "IRR", [f"=UnitEconomics!C{U['irr']}"], PCT)
srow("be_fee", "Break-even fee ($/kW-month)", [f"=UnitEconomics!C{U['be_fee']}"], USD1)
srow("yield", "Stabilised EBITDA yield on build cost", [f"=UnitEconomics!C{U['yield']}"], PCT)
rr += 1
hdr(wsS, rr, ["9. What it would take (Base supply, Targeted hub, 2030)", "Value", "", ""])
rr += 1
srow("lv_total", "Five-lever package (firm MW)", [f"=Levers!C{L['total']}"], NUM)
srow("lv_gapmin", "Gap to minimum reserve margin (MW)", [f"=Levers!C{L['gapmin']}"], NUM)
srow("lv_gaptgt", "Gap to target reserve margin (MW)", [f"=Levers!C{L['gaptgt']}"], NUM)
srow("lv_covtgt", "Coverage of the target-margin gap (x)", [f"=Levers!C{L['cov_tgt']}"], '0.00"x"')
srow("lv_cost", "Indicative package cost ($m/yr)", [f"=Levers!G{L['total']}"], NUM1)
srow("lv_rm", "Reserve margin in 2030 after the package", [f"=Levers!C{L['rm_after']}"], PCT)
srow("lv_self", "Share of package supplied by data centres themselves", [f"=Levers!C{L['self_share']}"], PCT)
srow("lv_costpct", "Hub's share of the target-margin gap", [f"=Levers!C{L['cost_pct']}"], PCT)

# =====================================================================
# SOURCES + CHECKS + README
# =====================================================================
wsR = wb.create_sheet("Sources")
title(wsR, "Source register and conflicts", "Figures read on 2026-10-19. Press and trade-press sources: confirm against EPRA / Kenya Power primary releases before publication.")
hdr(wsR, 4, ["ID", "Source", "Date", "URL", "What the model takes from it"])
for col, w in zip("ABCDE", [6, 70, 14, 90, 110]):
    wsR.column_dimensions[col].width = w
rr = 5
for sid, name, date, url, used in SOURCES:
    for ci, v in enumerate([sid, name, date, url, used], start=1):
        c = wsR.cell(rr, ci, v)
        c.font = f_norm
        c.alignment = Alignment(wrap_text=True, vertical="top")
    rr += 1
rr += 1
hdr(wsR, rr, ["ID", "Conflict between sources", "", "", "How the model resolves it"])
rr += 1
for cid, conflict, res in CONFLICTS:
    wsR.cell(rr, 1, cid).font = f_norm
    wsR.cell(rr, 2, conflict).font = f_norm
    wsR.cell(rr, 2).alignment = Alignment(wrap_text=True, vertical="top")
    wsR.cell(rr, 5, res).font = f_norm
    wsR.cell(rr, 5).alignment = Alignment(wrap_text=True, vertical="top")
    rr += 1

wsC = wb.create_sheet("Checks")
title(wsC, "Model integrity checks", "All checks should read TRUE.")
hdr(wsC, 4, ["Check", "Result"])
wsC.column_dimensions["A"].width = 100
wsC.column_dimensions["B"].width = 12
checks = [
    ("Base reserve margin reproduces the sourced 3.34% in 2026", f"=ABS(Grid!C{G[('Base','rm')]}-{I('rm0')})<0.000001"),
    ("Hub Targeted IT capacity in 2030 equals share x Africa 2030 capacity", f"=ABS(Hub!{Y2030}{H[('Targeted hub','it')]}-{I('share','E')}*{I('africa2030')})<0.000001"),
    ("Hub Status quo adds no incremental grid load", f"=Hub!{Y2030}{H[('Status quo','load')]}=0"),
    ("Campus NPV linear decomposition reproduces NPV", f"=ABS(UnitEconomics!C{U['chk']})<0.0001"),
    ("Reserve margin matrix, Base supply / Status quo equals Grid reserve margin in 2030", f"=ABS(Matrix!{M[('rm2030','Base','Status quo')]}-Grid!{Y2030}{G[('Base','rm')]})<0.000001"),
    ("Stress supply is never better than Base in 2030 (no hub)", f"=Grid!{Y2030}{G[('Stress','rm')]}<=Grid!{Y2030}{G[('Base','rm')]}"),
    ("Base supply is never better than Favourable in 2030 (no hub)", f"=Grid!{Y2030}{G[('Base','rm')]}<=Grid!{Y2030}{G[('Favourable','rm')]}"),
    ("Tenant all-in cost equals sum of components (Base)", f"=ABS(UnitEconomics!D{U['tco']}-UnitEconomics!D{U['fee']}-UnitEconomics!D{U['power']}-UnitEconomics!D{U['evening']})<0.000001"),
    ("Screen weights sum to 100%", "=ABS(Screen!G5-1)<0.000001"),
    ("Levers: package MW equals sum of lever rows", f"=ABS(Levers!C{L['total']}-SUM(Levers!C5:C{L['total']-1}))<0.000001"),
]
rr = 5
for lab, fm in checks:
    wsC.cell(rr, 1, lab).font = f_norm
    wsC.cell(rr, 2, fm).font = f_norm
    rr += 1
wsC.cell(rr + 1, 1, "All checks pass").font = f_bold
wsC.cell(rr + 1, 2, f"=AND(B5:B{rr-1})").font = f_bold

wsD = wb.create_sheet("README", 0)
title(wsD, "Kenya as East Africa's low-carbon industry hub - strategy model")
lines = [
    "Question: Should Kenya position itself as East Africa's hub for energy-intensive, low-carbon industry (data centres first), and what would it take?",
    "",
    "How to read this workbook",
    "  Summary         Headline outputs, all linked. Start here.",
    "  Inputs          Every assumption, tagged [S] sourced or [A] analyst assumption. Change blue cells; everything recalculates.",
    "  Grid            Peak demand versus firm, evening-capable supply, 2026-2032, in three supply cases (Stress / Base / Favourable).",
    "  Hub             Data-centre demand in three hub cases (Status quo / Targeted / Ambitious) and the size of the prize in 2030.",
    "  Matrix          Reserve margin, firm-MW gap and connectable load for every supply x hub combination.",
    "  UnitEconomics   Tenant cost per kW-month and an illustrative 44 MW campus (NPV, IRR, break-even fee, sensitivities).",
    "  Levers          Five levers sized against the 2030 firm-capacity gap, with indicative annual cost.",
    "  Screen          Judgement-based screen of which industries fit Kenya's evening-constrained grid.",
    "  Sources         Source register (S1-S27) and the conflicts between sources and how each is resolved.",
    "  Checks          Integrity checks; all should read TRUE.",
    "",
    "Colour code: blue = input; black = formula; green = link from another sheet; yellow fill = key lever.",
    "",
    "Limits you should know about",
    "  1. Several inputs come from press reports quoting Kenya Power / EPRA briefings, not primary datasets. Confirm against EPRA statistics before publishing.",
    "  2. Colocation pricing, utilisation, opex and the discount rate are analyst assumptions because no public Kenyan price data was available.",
    "  3. 'Firm' means available at the 6-10pm peak. Solar is therefore counted at 0%; storage at its rated MW for a 4-hour peak.",
    "  4. A negative reserve margin means unserved peak unless emergency capacity is procured; it sizes the response, it does not forecast blackouts.",
    "  5. The campus case is pre-tax and unlevered and meant to show what price or cost of capital a campus needs, not to value a specific project.",
]
for i, t in enumerate(lines, start=3):
    wsD.cell(i, 1, t).font = f_bold if t and not t.startswith(" ") and t.endswith(("workbook", "about")) else f_norm
wsD.column_dimensions["A"].width = 150

order = ["README", "Summary", "Inputs", "Grid", "Hub", "Matrix", "UnitEconomics", "Levers", "Screen", "Sources", "Checks"]
wb._sheets = [wb[n] for n in order]
wb.save(OUT)
print("saved", OUT)
